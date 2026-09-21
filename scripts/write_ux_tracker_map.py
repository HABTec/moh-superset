#!/usr/bin/env python3
"""Write the independent-plan → Excel tracker mapping (CSV + new sheet)."""

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ux_plan_tracker_map import OUT_OF_PLAN, ROWS

TRACKER = Path("/home/zelalem/Desktop/UX_UI/Week2_Task_Tracker(1).xlsx")
CSV = Path("/home/zelalem/Desktop/UX_UI/INDEPENDENT_PLAN_TRACKER_MAP.csv")
SHEET = "Independent plan map"

HEADERS = [
    "Plan ID",
    "Plan task",
    "Wave",
    "Tracker Task ID",
    "Finding ID",
    "Stream",
    "Severity",
    "Tracker status",
    "Coverage",
    "Tracker short title",
    "What this plan does",
    "What stays on the tracker",
]

NAVY = "1B365D"
TEAL = "0E6B7A"
WHITE = "FFFFFF"
PALE = "EEF3F7"
COVER = {
    "CLOSES": "1E7A4C",
    "PARTIAL": "8F5D00",
    "VERIFY": "0E6B7A",
    "NEW": "5B4B8A",
    "OUT": "4A5568",
}


def write_csv() -> None:
    import csv

    with CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        w.writerows(ROWS)
        w.writerow([])
        w.writerow(
            [
                "OUT OF THIS PLAN",
                "Tracker Task ID",
                "Finding",
                "Severity",
                "Title",
                "Why not in the PDF",
            ]
        )
        w.writerows(OUT_OF_PLAN)


def style_header(ws, row: int, ncols: int) -> None:
    fill = PatternFill("solid", fgColor=NAVY)
    font = Font(name="Calibri", bold=True, color=WHITE, size=11)
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def write_sheet(wb: openpyxl.Workbook) -> None:
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET, 0)

    ws["A1"] = "Independent implementation plan  →  Week 2 task tracker"
    ws["A1"].font = Font(name="Calibri", bold=True, size=16, color=NAVY)
    ws.merge_cells("A1:L1")
    ws["A2"] = (
        "Source: Week2_Task_Tracker(1).xlsx Tasks sheet  ·  Plan: "
        "INDEPENDENT_IMPLEMENTATION_PLAN.pdf  ·  20 September 2026.  "
        "CLOSES = mark the A-row done after this plan task.  "
        "PARTIAL = do the no-wait slice; tracker row stays open.  "
        "VERIFY = tracker already DONE; confirm in UI.  "
        "NEW = not in the tracker (Yengwe TV).  "
        "OUT = skip."
    )
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:L2")
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 48

    start = 4
    for c, h in enumerate(HEADERS, 1):
        ws.cell(start, c, h)
    style_header(ws, start, len(HEADERS))

    thin = Border(
        left=Side(style="thin", color="D6DEE8"),
        right=Side(style="thin", color="D6DEE8"),
        top=Side(style="thin", color="D6DEE8"),
        bottom=Side(style="thin", color="D6DEE8"),
    )
    wrap = Alignment(wrap_text=True, vertical="top")
    body = Font(name="Calibri", size=10)
    bold = Font(name="Calibri", size=10, bold=True, color=WHITE)

    for i, row in enumerate(ROWS):
        r = start + 1 + i
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c, val)
            cell.alignment = wrap
            cell.border = thin
            cell.font = body
            if i % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=PALE)
        cov = row[8]
        cov_cell = ws.cell(r, 9)
        cov_cell.fill = PatternFill("solid", fgColor=COVER.get(cov, "4A5568"))
        cov_cell.font = bold
        cov_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r].height = 48

    last = start + len(ROWS)

    out_r = last + 3
    ws.cell(out_r, 1, "Tracker rows this plan does not start")
    ws.cell(out_r, 1).font = Font(name="Calibri", bold=True, size=13, color=NAVY)
    ws.merge_cells(start_row=out_r, start_column=1, end_row=out_r, end_column=6)
    hdr_r = out_r + 1
    out_headers = [
        "Tracker Task ID",
        "Finding ID",
        "Severity",
        "Title",
        "Why it is out of the independent PDF",
    ]
    for c, h in enumerate(out_headers, 1):
        ws.cell(hdr_r, c, h)
    style_header(ws, hdr_r, 5)
    for i, row in enumerate(OUT_OF_PLAN):
        r = hdr_r + 1 + i
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c, val)
            cell.alignment = wrap
            cell.border = thin
            cell.font = body
            if i % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=PALE)
        ws.row_dimensions[r].height = 28

    widths = [10, 42, 8, 16, 12, 10, 12, 14, 12, 42, 52, 48]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(HEADERS))}{last}"
    ws.sheet_properties.tabColor = TEAL


def main() -> None:
    write_csv()
    wb = openpyxl.load_workbook(TRACKER)
    write_sheet(wb)
    wb.save(TRACKER)
    print("Wrote sheet", SHEET, "on", TRACKER)
    print("Wrote", CSV)


if __name__ == "__main__":
    main()
